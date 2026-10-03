import base64
import hashlib
import json
import re
import secrets
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "密碼.xlsx"
OUTPUT = ROOT / "src" / "data" / "credentials.generated.ts"


def encode_password(student_id: str, password: str) -> dict[str, str]:
    salt = secrets.token_hex(8)
    key = hashlib.sha256(f"{student_id}:{salt}".encode("utf-8")).digest()
    password_bytes = password.encode("utf-8")
    encrypted = bytes(value ^ key[index % len(key)] for index, value in enumerate(password_bytes))
    return {"salt": salt, "value": base64.b64encode(encrypted).decode("ascii")}


def main() -> None:
    if not WORKBOOK.is_file():
        raise FileNotFoundError("找不到密碼.xlsx，請將原始資料放在專案根目錄。")

    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = next(rows, None)
        columns = {
            str(value).strip(): index
            for index, value in enumerate(headers or ())
            if value is not None
        }
        if "帳號" not in columns or "密碼" not in columns:
            raise ValueError("Excel 必須有「帳號」及「密碼」欄位。")

        credentials = {}
        for row_number, row in enumerate(rows, start=2):
            account_value = row[columns["帳號"]] if columns["帳號"] < len(row) else None
            password_value = row[columns["密碼"]] if columns["密碼"] < len(row) else None
            if account_value is None and password_value is None:
                continue

            account = str(account_value).strip() if account_value is not None else ""
            password = str(password_value).strip() if password_value is not None else ""
            match = re.fullmatch(r"s(\d+)", account, re.IGNORECASE)
            if not match or not password:
                raise ValueError(f"Excel 第 {row_number} 列帳號或密碼格式不正確。")

            student_id = match.group(1)
            if student_id in credentials:
                raise ValueError(f"Excel 第 {row_number} 列學號重複。")
            credentials[student_id] = encode_password(student_id, password)
    finally:
        workbook.close()

    source = (
        "// Generated from the local workbook. The values are obfuscated, not encrypted.\n"
        "export type EncodedCredential = { salt: string; value: string }\n\n"
        "export const encodedCredentials: Record<string, EncodedCredential> = "
        f"{json.dumps(credentials, ensure_ascii=True, indent=2)}\n"
    )
    OUTPUT.write_text(source, encoding="utf-8")
    print(f"已產生 {len(credentials)} 筆混淆資料：{OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()