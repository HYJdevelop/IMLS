import os
import re
from pathlib import Path

from flask import Flask, render_template, request, send_from_directory
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
FRONTEND_DIST = BASE_DIR / "dist"
configured_workbook = os.environ.get("CREDENTIALS_FILE")
WORKBOOK_PATH = Path(configured_workbook).expanduser() if configured_workbook else BASE_DIR / "密碼.xlsx"
if not WORKBOOK_PATH.is_absolute():
    WORKBOOK_PATH = BASE_DIR / WORKBOOK_PATH

app = Flask(
    __name__,
    template_folder=str(BASE_DIR / "templates"),
    static_folder=str(FRONTEND_DIST / "assets"),
    static_url_path="/assets",
)
app.config["MAX_CONTENT_LENGTH"] = 4 * 1024


def load_credentials():
    if not WORKBOOK_PATH.is_file():
        raise FileNotFoundError(
            "找不到伺服器帳密檔。請設定 CREDENTIALS_FILE 指向受保護的密碼.xlsx。"
        )

    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValueError("帳密資料表是空的。")

        columns = {
            str(value).strip(): index
            for index, value in enumerate(headers)
            if value is not None
        }
        if "帳號" not in columns or "密碼" not in columns:
            raise ValueError("帳密資料表必須包含「帳號」和「密碼」欄位。")

        credentials = {}
        for row_number, row in enumerate(rows, start=2):
            account_index = columns["帳號"]
            password_index = columns["密碼"]
            account_value = row[account_index] if account_index < len(row) else None
            password_value = row[password_index] if password_index < len(row) else None
            if account_value is None and password_value is None:
                continue

            account = str(account_value).strip() if account_value is not None else ""
            password = str(password_value).strip() if password_value is not None else ""
            match = re.fullmatch(r"s(\d+)", account, re.IGNORECASE)
            if not match or not password:
                raise ValueError(f"帳密資料表第 {row_number} 列的帳號或密碼格式不正確。")

            student_id = match.group(1)
            if student_id in credentials:
                raise ValueError(f"帳密資料表有重複學號（第 {row_number} 列）。")
            credentials[student_id] = {"account": account, "password": password}

        return credentials
    finally:
        workbook.close()


CREDENTIALS = load_credentials()


@app.after_request
def set_security_headers(response):
    response.headers["Cache-Control"] = "no-store"
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; style-src 'self'; script-src 'self'; "
        "img-src 'self' data: https://www.hyjdevelop.com; form-action 'self'; "
        "base-uri 'self'; frame-ancestors 'none'"
    )
    return response


@app.get("/")
def index():
    if not (FRONTEND_DIST / "index.html").is_file():
        return "請先執行 npm run build。", 503
    return send_from_directory(FRONTEND_DIST, "index.html")


@app.get("/result.css")
def result_css():
    return send_from_directory(FRONTEND_DIST, "result.css")


@app.post("/lookup")
def lookup():
    student_id = request.form.get("student_id", "").strip()
    if not re.fullmatch(r"\d+", student_id):
        return render_template("result.html", error="請輸入不含 s 的數字學號。"), 400

    credential = CREDENTIALS.get(student_id)
    if credential is None:
        return render_template("result.html", error="查無此學號，請確認輸入是否正確。"), 404

    return render_template("result.html", credential=credential)


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", "5000")), debug=False)
